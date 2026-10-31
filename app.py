"""
app.py
======
Buscador de Planes de Cuidados de Enfermería.

El Excel institucional vive dentro del proyecto (carpeta ``data/``) y actúa
como base de datos estática: no hay carga de archivos. Para actualizar la
información basta con reemplazar ``data/planes_enfermeria.xlsx``.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

import openpyxl
import pandas as pd
import streamlit as st

# ---------------------------------------------------------------------------
# Configuración general
# ---------------------------------------------------------------------------

st.set_page_config(
    page_title="Buscador de Planes de Enfermería",
    page_icon="🩺",
    layout="wide",
)

RUTA_EXCEL = Path(__file__).parent / "data" / "planes_enfermeria.xlsx"

PLACEHOLDER = "No especificado en el documento original"

PALABRAS_CLAVE_ENCABEZADO = [
    "diagnostico", "riesgo", "dominio", "hallazgo", "resultado",
    "intervencion", "actividad", "evaluacion", "prevenir", "cuidado",
    "dx", "enferm",
]
MAX_FILAS_ESCANEO_ENCABEZADO = 25
MIN_COINCIDENCIAS_ENCABEZADO = 2

SINONIMOS_COLUMNAS = {
    "riesgo": ["riesgo a prevenir", "riesgo"],
    "dominio": ["dominio"],
    "hallazgos": ["hallazgo"],
    "diagnostico": ["diagnostico"],
    "diagnostico_alt": ["dx"],
    "resultado": ["resultado esperado", "resultado"],
    "intervencion": ["intervencion"],
    "actividad": ["actividad"],
    "evaluacion": ["evaluacion"],
}


# ---------------------------------------------------------------------------
# Utilidades de texto
# ---------------------------------------------------------------------------

def _normalizar(texto: object) -> str:
    """Texto en minúsculas y sin tildes, seguro ante NaN."""
    if texto is None or (isinstance(texto, float) and pd.isna(texto)):
        return ""
    s = str(texto).strip().lower()
    for origen, destino in {"á": "a", "é": "e", "í": "i", "ó": "o", "ú": "u", "ñ": "n"}.items():
        s = s.replace(origen, destino)
    return s


def _normalizar_espacios(texto: str) -> str:
    """Colapsa saltos de línea y espacios múltiples en un solo espacio."""
    return " ".join(texto.split())


def _md(texto: str) -> str:
    """Prepara el texto de una celda para markdown: conserva los saltos de
    línea (que markdown ignoraría) y elimina tabs sobrantes."""
    lineas = [l.strip() for l in texto.replace("\t", " ").splitlines()]
    return "  \n".join(l for l in lineas if l)


def _es_valor_vacio(valor: object) -> bool:
    if valor is None:
        return True
    if isinstance(valor, float) and pd.isna(valor):
        return True
    if isinstance(valor, str) and valor.strip() == "":
        return True
    return False


def _texto_limpio(valor: object) -> Optional[str]:
    """Devuelve el valor como texto limpio (o None si está vacío)."""
    if _es_valor_vacio(valor):
        return None
    return str(valor).strip().replace("<", "&lt;").replace(">", "&gt;")


# ---------------------------------------------------------------------------
# Detección de encabezado y columnas
# ---------------------------------------------------------------------------

def detectar_fila_encabezado(df_crudo: pd.DataFrame) -> Optional[int]:
    n_filas = min(MAX_FILAS_ESCANEO_ENCABEZADO, len(df_crudo))
    mejor_fila, mejor_puntaje = None, 0

    for i in range(n_filas):
        celdas = [_normalizar(v) for v in df_crudo.iloc[i].tolist()]
        puntaje = sum(1 for c in celdas if c and any(k in c for k in PALABRAS_CLAVE_ENCABEZADO))
        no_vacias = sum(1 for c in celdas if c)
        if puntaje >= MIN_COINCIDENCIAS_ENCABEZADO and no_vacias >= puntaje:
            if puntaje > mejor_puntaje:
                mejor_puntaje, mejor_fila = puntaje, i

    return mejor_fila


def encontrar_columna(columnas: list[str], claves: list[str]) -> Optional[str]:
    for col in columnas:
        n = _normalizar(col)
        if any(clave in n for clave in claves):
            return col
    return None


def identificar_columnas_clave(columnas: list[str]) -> dict[str, Optional[str]]:
    mapa = {
        nombre: encontrar_columna(columnas, SINONIMOS_COLUMNAS[nombre])
        for nombre in (
            "riesgo", "dominio", "hallazgos", "resultado",
            "intervencion", "actividad", "evaluacion",
        )
    }
    diag = encontrar_columna(columnas, SINONIMOS_COLUMNAS["diagnostico"])
    if diag is None:
        diag = encontrar_columna(columnas, SINONIMOS_COLUMNAS["diagnostico_alt"])
    mapa["diagnostico"] = diag
    return mapa


# ---------------------------------------------------------------------------
# Estructuras de datos
# ---------------------------------------------------------------------------

@dataclass
class Intervencion:
    texto: Optional[str]
    actividades: list[str] = field(default_factory=list)


@dataclass
class CasoEnfermeria:
    hoja: str
    diagnostico: Optional[str]
    riesgo: Optional[str]
    dominio: Optional[str]
    resultado: Optional[str]
    evaluacion: Optional[str]
    hallazgos: list[str] = field(default_factory=list)
    intervenciones: list[Intervencion] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Mini-ETL
# ---------------------------------------------------------------------------

def _leer_hoja_con_combinadas(ws) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Devuelve dos tablas del mismo tamaño:
      - crudo: solo la celda superior-izquierda de cada celda combinada
        (sirve para saber dónde EMPIEZA cada diagnóstico / intervención).
      - lleno: el valor de la celda combinada repetido en todas las filas
        que abarca (sirve para leer hallazgos, dominio, resultado, etc.)."""
    filas = [list(r) for r in ws.iter_rows(values_only=True)]
    crudo = pd.DataFrame(filas, dtype=object)
    lleno = crudo.copy()
    for rng in ws.merged_cells.ranges:
        valor = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row - 1, rng.max_row):
            for c in range(rng.min_col - 1, rng.max_col):
                if r < lleno.shape[0] and c < lleno.shape[1]:
                    lleno.iat[r, c] = valor
    return crudo, lleno


def _procesar_hoja(nombre_hoja: str, df_crudo: pd.DataFrame, df_lleno: pd.DataFrame) -> list[CasoEnfermeria]:
    if df_crudo.empty:
        return []

    idx_encabezado = detectar_fila_encabezado(df_crudo)
    if idx_encabezado is None:
        return []

    nombres = [
        str(c).strip() if pd.notna(c) else f"Columna_{i}"
        for i, c in enumerate(df_crudo.iloc[idx_encabezado].tolist())
    ]
    df_hoja = df_crudo.iloc[idx_encabezado + 1:].copy()
    df_full = df_lleno.iloc[idx_encabezado + 1:].copy()
    df_hoja.columns = nombres
    df_full.columns = nombres

    # Se descartan solo las filas totalmente vacías en la hoja original.
    con_datos = df_hoja.notna().any(axis=1)
    df_hoja = df_hoja[con_datos].reset_index(drop=True)
    df_full = df_full[con_datos].reset_index(drop=True)
    if df_hoja.empty:
        return []

    columnas = identificar_columnas_clave(list(df_hoja.columns))
    col_diag = columnas["diagnostico"]
    if col_diag is None:
        return []

    def valor_para(col: Optional[str], pos: int) -> Optional[str]:
        return None if col is None else _texto_limpio(df_full.at[pos, col])

    # Cada diagnóstico no vacío abre un caso; cada intervención no vacía,
    # una intervención dentro del caso.
    bloque_diag = df_hoja[col_diag].apply(lambda v: not _es_valor_vacio(v)).cumsum()
    col_interv = columnas["intervencion"]
    if col_interv is not None:
        bloque_interv = df_hoja[col_interv].apply(lambda v: not _es_valor_vacio(v)).cumsum()
    else:
        bloque_interv = bloque_diag

    df_trabajo = pd.DataFrame({"_bloque_diag": bloque_diag, "_bloque_interv": bloque_interv})
    df_trabajo["_pos"] = range(len(df_trabajo))
    df_trabajo = df_trabajo[df_trabajo["_bloque_diag"] >= 1]

    col_hallazgos = columnas["hallazgos"]
    col_actividad = columnas["actividad"]

    casos: list[CasoEnfermeria] = []
    for _, grupo_diag in df_trabajo.groupby("_bloque_diag", sort=True):
        primera = int(grupo_diag["_pos"].iloc[0])
        posiciones = [int(p) for p in grupo_diag["_pos"]]

        # Los hallazgos suelen estar en una celda combinada que abarca varios
        # diagnósticos: todos ellos comparten esos hallazgos.
        hallazgos: list[str] = []
        if col_hallazgos is not None:
            for p in posiciones:
                t = _texto_limpio(df_full.at[p, col_hallazgos])
                if t and t not in hallazgos:
                    hallazgos.append(t)

        intervenciones: list[Intervencion] = []
        for _, sub in grupo_diag.groupby("_bloque_interv", sort=True):
            pos_sub = [int(p) for p in sub["_pos"]]
            texto = valor_para(col_interv, pos_sub[0]) if col_interv else None

            actividades: list[str] = []
            if col_actividad is not None:
                for p in pos_sub:
                    t = _texto_limpio(df_full.at[p, col_actividad])
                    if t and t not in actividades:
                        actividades.append(t)

            if texto or actividades:
                intervenciones.append(Intervencion(texto=texto, actividades=actividades))

        casos.append(CasoEnfermeria(
            hoja=nombre_hoja.strip(),
            diagnostico=valor_para(col_diag, primera),
            riesgo=valor_para(columnas["riesgo"], primera),
            dominio=valor_para(columnas["dominio"], primera),
            resultado=valor_para(columnas["resultado"], primera),
            evaluacion=valor_para(columnas["evaluacion"], primera),
            hallazgos=hallazgos,
            intervenciones=intervenciones,
        ))

    return casos


@st.cache_resource(show_spinner="Cargando base de datos...")
def cargar_base_de_datos(ruta: str, mtime: float) -> list[CasoEnfermeria]:
    """Lee el Excel interno. ``mtime`` invalida la caché si el archivo cambia."""
    libro = openpyxl.load_workbook(ruta, data_only=True)
    casos: list[CasoEnfermeria] = []
    for ws in libro.worksheets:
        try:
            crudo, lleno = _leer_hoja_con_combinadas(ws)
            casos.extend(_procesar_hoja(ws.title, crudo, lleno))
        except Exception:
            continue
    return casos


# ---------------------------------------------------------------------------
# Renderizado UI
# ---------------------------------------------------------------------------

def _mostrar_campo(etiqueta: str, valor: Optional[str]) -> None:
    st.markdown(f"**{etiqueta}:** {_md(valor) if valor else PLACEHOLDER}")


def _mostrar_resultado(valor: Optional[str]) -> None:
    _mostrar_campo("Resultado esperado (NOC)", valor)


def _renderizar_intervenciones(intervenciones: list[Intervencion]) -> None:
    st.markdown("**Intervenciones / Actividades:**")
    if not intervenciones:
        st.markdown(f"- {PLACEHOLDER}")
        return

    for interv in intervenciones:
        titulo = _normalizar_espacios(interv.texto) if interv.texto else PLACEHOLDER
        n = len(interv.actividades)
        sufijo = f" · {n} actividad{'es' if n != 1 else ''}" if n else ""
        with st.expander(f"🩹 {titulo}{sufijo}", expanded=False):
            if interv.actividades:
                for act in interv.actividades:
                    st.markdown(f"- {_md(act)}")
            else:
                st.markdown(PLACEHOLDER)


def _cajas_por_diagnostico(casos_hoja: list[CasoEnfermeria]) -> list[tuple[str, CasoEnfermeria]]:
    """Una cajita independiente por cada diagnóstico de enfermería. El título
    son sus hallazgos; si varios diagnósticos comparten los mismos hallazgos,
    se numera para distinguirlos."""
    cajas: list[tuple[str, CasoEnfermeria]] = []
    for caso in casos_hoja:
        if caso.hallazgos:
            titulo = " ".join(_normalizar_espacios(h) for h in caso.hallazgos)
        else:
            titulo = f"Sin hallazgo específico — {_normalizar_espacios(caso.diagnostico or PLACEHOLDER)}"
        if len(titulo) > 140:
            titulo = titulo[:137] + "..."
        cajas.append((titulo, caso))

    totales: dict[str, int] = {}
    for titulo, _ in cajas:
        totales[titulo] = totales.get(titulo, 0) + 1

    vistos: dict[str, int] = {}
    finales: list[tuple[str, CasoEnfermeria]] = []
    for titulo, caso in cajas:
        if totales[titulo] > 1:
            vistos[titulo] = vistos.get(titulo, 0) + 1
            titulo_final = f"{titulo}  ·  diagnóstico {vistos[titulo]} de {totales[titulo]}"
        else:
            titulo_final = titulo
        finales.append((titulo_final, caso))
    return finales


# ---------------------------------------------------------------------------
# Estilos
# ---------------------------------------------------------------------------

COLOR_INICIO = "#B37FE0"
COLOR_DX_MEDICO = "#1FAE6B"
COLOR_HALLAZGOS = "#FF3FA4"
COLOR_DOMINIO = "#3FA7FF"
COLOR_DIAGNOSTICO = "#FF7FC0"


def _inyectar_estilos_flujo() -> None:
    st.markdown(
        """
        <style>
        .pill-paso {
            display: inline-block;
            padding: 8px 18px;
            border-radius: 999px;
            color: white;
            font-weight: 700;
            margin: 6px 0 10px 0;
            font-size: 0.95rem;
            white-space: normal;
            overflow-wrap: break-word;
            max-width: 100%;
        }
        .subtitulo-seccion {
            font-size: 0.8rem;
            font-weight: 800;
            letter-spacing: 0.06em;
            color: #6B5B95;
            margin: 10px 0 0 2px;
        }
        .caja-hallazgos {
            background: #FFE3F3;
            border-left: 6px solid #FF3FA4;
            border-radius: 10px;
            padding: 14px 18px;
            margin-bottom: 8px;
        }
        div[data-testid="stExpander"] {
            border: 1px solid rgba(0,0,0,0.08) !important;
            border-radius: 12px !important;
            margin-bottom: 8px;
            overflow: hidden;
        }
        div[data-testid="stExpander"] summary {
            font-weight: 600 !important;
            background: #F3F0FB;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )


def _pill(texto: str, color: str) -> None:
    st.markdown(f'<span class="pill-paso" style="background:{color};">{texto}</span>', unsafe_allow_html=True)


# ---------------------------------------------------------------------------
# Vistas
# ---------------------------------------------------------------------------

def _vista_guiada_por_diagnostico_medico(casos: list[CasoEnfermeria]) -> None:
    hojas = sorted({c.hoja for c in casos}, key=_normalizar)

    _pill("🟣 Inicio", COLOR_INICIO)
    st.write("")
    _pill("1️⃣ Diagnóstico médico", COLOR_DX_MEDICO)
    hoja = st.selectbox(
        "Selecciona la patología / diagnóstico médico:",
        options=hojas,
        index=None,
        placeholder="Escribe o selecciona un diagnóstico médico...",
        label_visibility="collapsed",
        key="selector_dx_medico",
    )
    if hoja is None:
        st.caption("👆 Elige un diagnóstico médico para continuar.")
        return

    casos_hoja = [c for c in casos if c.hoja == hoja]
    _pill(
        "2️⃣ Hallazgos → Diagnóstico de enfermería — despliega una cajita "
        "para ver su diagnóstico, resultado e intervenciones",
        COLOR_HALLAZGOS,
    )

    cajas = _cajas_por_diagnostico(casos_hoja)
    if not cajas:
        st.markdown(f'<div class="caja-hallazgos">{PLACEHOLDER}</div>', unsafe_allow_html=True)
        return

    for titulo, caso in cajas:
        with st.expander(f"🔎 {titulo}", expanded=False):
            for h in caso.hallazgos:
                st.markdown(f"- {_md(h)}")
            if caso.hallazgos:
                st.write("")

            _pill(f"🧭 Dominio: {caso.dominio or PLACEHOLDER}", COLOR_DOMINIO)
            st.markdown('<div class="subtitulo-seccion">DIAGNÓSTICO DE ENFERMERÍA</div>', unsafe_allow_html=True)
            _pill(f"🩺 {caso.diagnostico or PLACEHOLDER}", COLOR_DIAGNOSTICO)
            _mostrar_resultado(caso.resultado)
            _renderizar_intervenciones(caso.intervenciones)


# ---------------------------------------------------------------------------
# Interfaz principal
# ---------------------------------------------------------------------------

def main() -> None:
    _inyectar_estilos_flujo()
    st.title("🩺 Buscador de Planes de Cuidados de Enfermería")
    st.caption("Selecciona un diagnóstico médico para consultar sus planes de cuidado de enfermería.")

    if not RUTA_EXCEL.exists():
        st.error(f"No se encontró la base de datos en: {RUTA_EXCEL}")
        return

    try:
        casos = cargar_base_de_datos(str(RUTA_EXCEL), RUTA_EXCEL.stat().st_mtime)
    except Exception as e:
        st.error(f"No se pudo leer la base de datos. Detalle: {e}")
        return

    if not casos:
        st.error("No se pudo extraer ningún caso de enfermería de la base de datos.")
        return

    _vista_guiada_por_diagnostico_medico(casos)


if __name__ == "__main__":
    main()
